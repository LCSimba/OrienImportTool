"""Low-level parser for Orien Tactics single-sheet exports.

Validates the metadata header, builds a ``column_name -> index`` map from the
machine-readable row, and yields one dict per data row with empty cells
normalised to ``""``. Semantic grouping into the canonical domain model is the
:mod:`normalizer` module's job.
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook import Workbook

from orien_import_tool.importers.orien import schema as S


class OrienParseError(ValueError):
    """Raised when an Orien export violates its expected structure."""


@dataclass(frozen=True)
class ParsedHeader:
    """Metadata extracted from rows 1-7 of the data sheet.

    ``structure_revision_token`` is the UUID-style snapshot id from row 7's
    ``revision`` key. The integer revision number is per-row on the data sheet
    (column ``structureRevision``) and is captured by the normalizer onto
    :class:`Equipment.structure_revision`.
    """

    location_token: str
    location_description: str
    language: str
    export_timestamp: datetime | None
    structure_revision_token: str
    component_library: bool


@dataclass
class ParsedExport:
    """A parsed Orien export ready for normalization."""

    source_path: Path
    header: ParsedHeader
    column_index: dict[str, int]
    rows: list[dict[str, Any]] = field(default_factory=list)


def parse_workbook(path: Path) -> ParsedExport:
    """Parse an Orien single-sheet export from ``path``.

    Raises :class:`OrienParseError` for any structural drift from the contract
    documented in ``tests/fixtures/orien/SCHEMA.md``.
    """

    if not path.exists():
        raise OrienParseError(f"Orien export not found: {path}")
    wb: Workbook = openpyxl.load_workbook(path, data_only=True)
    _validate_sheet_names(wb, path)

    ws = wb[S.SHEET_DATA]
    grid = list(ws.iter_rows(values_only=True))
    if len(grid) <= S.FIRST_DATA_ROW:
        raise OrienParseError(f"{path.name}: no data rows after header")

    _validate_header_markers(grid, path)
    header = _extract_header(grid, path)
    column_index = _extract_columns(grid, path)
    rows = list(_iter_data_rows(grid, column_index))
    return ParsedExport(
        source_path=path,
        header=header,
        column_index=column_index,
        rows=rows,
    )


# --- Internal helpers -----------------------------------------------------------------


def _validate_sheet_names(wb: Workbook, path: Path) -> None:
    found = tuple(wb.sheetnames)
    if found != S.EXPECTED_SHEET_NAMES:
        raise OrienParseError(
            f"{path.name}: unexpected sheet names {found!r}, expected {S.EXPECTED_SHEET_NAMES!r}"
        )


def _validate_header_markers(grid: list[tuple[Any, ...]], path: Path) -> None:
    export_marker = _cell(grid, S.ROW_EXPORT_MARKER, 0)
    if export_marker != S.EXPECTED_EXPORT_MARKER:
        raise OrienParseError(
            f"{path.name}: expected {S.EXPECTED_EXPORT_MARKER!r} on row "
            f"{S.ROW_EXPORT_MARKER + 1}, got {export_marker!r}"
        )
    sheet_type = _cell(grid, S.ROW_SHEET_TYPE_MARKER, 0)
    if sheet_type != S.EXPECTED_SHEET_TYPE_MARKER:
        raise OrienParseError(
            f"{path.name}: expected {S.EXPECTED_SHEET_TYPE_MARKER!r} on row "
            f"{S.ROW_SHEET_TYPE_MARKER + 1}, got {sheet_type!r}"
        )
    tactics_marker = _cell(grid, S.ROW_TACTICS_MARKER, 0)
    if tactics_marker != S.EXPECTED_TACTICS_MARKER:
        raise OrienParseError(
            f"{path.name}: expected {S.EXPECTED_TACTICS_MARKER!r} on row "
            f"{S.ROW_TACTICS_MARKER + 1}, got {tactics_marker!r}"
        )


def _extract_header(grid: list[tuple[Any, ...]], path: Path) -> ParsedHeader:
    location_row = grid[S.ROW_LOCATION]
    location_token = _str(location_row[0])
    location_description = _str(location_row[1]) if len(location_row) > 1 else ""
    language = _str(location_row[2]) if len(location_row) > 2 else ""
    if not location_token:
        raise OrienParseError(f"{path.name}: missing locationToken on row {S.ROW_LOCATION + 1}")
    if not location_description:
        raise OrienParseError(
            f"{path.name}: missing locationDescription on row {S.ROW_LOCATION + 1}"
        )

    tactics_row = grid[S.ROW_TACTICS_MARKER]
    pairs = _key_value_pairs(tactics_row[2:])
    location_token_check = pairs.get("location", "")
    if location_token_check and location_token_check != location_token:
        raise OrienParseError(
            f"{path.name}: locationToken mismatch — row {S.ROW_LOCATION + 1} has "
            f"{location_token!r} but row {S.ROW_TACTICS_MARKER + 1} has "
            f"{location_token_check!r}"
        )

    export_timestamp: datetime | None = None
    if len(tactics_row) > 1 and isinstance(tactics_row[1], datetime):
        export_timestamp = tactics_row[1]

    structure_revision_token = pairs.get("revision", "")
    component_library = pairs.get("componentLibrary", "").lower() == "true"

    return ParsedHeader(
        location_token=location_token,
        location_description=location_description,
        language=language or "en",
        export_timestamp=export_timestamp,
        structure_revision_token=structure_revision_token,
        component_library=component_library,
    )


def _extract_columns(grid: list[tuple[Any, ...]], path: Path) -> dict[str, int]:
    header_row = grid[S.ROW_MACHINE_HEADERS]
    columns: dict[str, int] = {}
    for idx, value in enumerate(header_row):
        name = _str(value)
        if not name:
            continue
        if name in columns:
            raise OrienParseError(
                f"{path.name}: duplicate column name {name!r} at indices {columns[name]} and {idx}"
            )
        columns[name] = idx
    missing = S.REQUIRED_COLUMNS - set(columns)
    if missing:
        raise OrienParseError(f"{path.name}: missing required columns {sorted(missing)}")
    return columns


def _iter_data_rows(
    grid: list[tuple[Any, ...]],
    column_index: dict[str, int],
) -> Iterator[dict[str, Any]]:
    for row in grid[S.FIRST_DATA_ROW :]:
        if _is_empty_row(row):
            continue
        record: dict[str, Any] = {}
        for name, idx in column_index.items():
            value = row[idx] if idx < len(row) else None
            record[name] = value if value is not None else ""
        yield record


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def _cell(grid: list[tuple[Any, ...]], row: int, col: int) -> Any:
    if row >= len(grid):
        return None
    line = grid[row]
    if col >= len(line):
        return None
    return line[col]


def _str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _key_value_pairs(values: Iterable[Any]) -> dict[str, str]:
    """Pair adjacent cells into ``{key: value}`` for the row 7 marker layout."""
    cleaned = [_str(v) for v in values]
    pairs: dict[str, str] = {}
    i = 0
    while i + 1 < len(cleaned):
        key = cleaned[i]
        if key:
            pairs[key] = cleaned[i + 1]
        i += 2
    return pairs
